import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_BG = RGBColor(248, 250, 252)         # #F8FAFC Slate 50
    COLOR_NAVY = RGBColor(15, 23, 42)          # #0F172A Slate 900
    COLOR_DARK = RGBColor(30, 41, 59)          # #1E293B Slate 800
    COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B Slate 500
    COLOR_BODY = RGBColor(51, 65, 85)          # #334155 Slate 700
    COLOR_CARD_BG = RGBColor(255, 255, 255)    # #FFFFFF White
    COLOR_CARD_BORDER = RGBColor(226, 232, 240)# #E2E8F0 Slate 200
    COLOR_BLUE = RGBColor(37, 99, 235)         # #2563EB Blue 600
    COLOR_RED = RGBColor(220, 38, 38)          # #DC2626 Red 600
    COLOR_GREEN = RGBColor(22, 163, 74)        # #16A34A Green 600
    COLOR_AMBER = RGBColor(217, 119, 6)        # #D97706 Amber 600

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG
        # Send to back if needed, but since it's added first, it's already in the background

    def add_header(slide, slide_num, total_slides, section_num, section_title, subtitle):
        # Top header row: course tag and slide counter
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.35))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = "23CSE304  |  EMBEDDED SYSTEMS"
        r1.font.size = Pt(11)
        r1.font.bold = True
        r1.font.color.rgb = COLOR_MUTED

        r2 = p.add_run()
        r2.text = f"                                                                                                                  {slide_num} / {total_slides}"
        r2.font.size = Pt(11)
        r2.font.bold = True
        r2.font.color.rgb = COLOR_MUTED

        # Main Section Title & Subtitle
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.9))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = f"{section_num} • {section_title.upper()}"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY

        p_sub = tf_title.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = COLOR_MUTED

    def add_card(slide, left, top, width, height, title, items, title_color=COLOR_NAVY, tag=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.2)

        tx_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), width - Inches(0.6), height - Inches(0.5))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Title / Tag
        p = tf.paragraphs[0]
        if tag:
            rtag = p.add_run()
            rtag.text = tag.upper() + "\n"
            rtag.font.size = Pt(10)
            rtag.font.bold = True
            rtag.font.color.rgb = COLOR_BLUE

        rtitle = p.add_run()
        rtitle.text = title
        rtitle.font.size = Pt(15)
        rtitle.font.bold = True
        rtitle.font.color.rgb = title_color
        p.space_after = Pt(10)

        # Content items
        for bold_prefix, text in items:
            p_item = tf.add_paragraph()
            p_item.space_after = Pt(7)
            if bold_prefix:
                rbold = p_item.add_run()
                rbold.text = "• " + bold_prefix + ": "
                rbold.font.size = Pt(12)
                rbold.font.bold = True
                rbold.font.color.rgb = COLOR_DARK

                rtext = p_item.add_run()
                rtext.text = text
                rtext.font.size = Pt(12)
                rtext.font.color.rgb = COLOR_BODY
            else:
                rtext = p_item.add_run()
                rtext.text = "• " + text
                rtext.font.size = Pt(12)
                rtext.font.color.rgb = COLOR_BODY

    blank_layout = prs.slide_layouts[6]

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(4.8))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "23CSE304 — EMBEDDED SYSTEMS"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_BLUE
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "SMART AMBULANCE TRAFFIC PRIORITY\nAND JUNCTION MANAGEMENT SYSTEM"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NAVY
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "An Intelligent, Interrupt-Driven Controller with Anti-Starvation Scheduling on STM32C031C6"
    p2.font.size = Pt(15)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "TEAM MEMBERS:"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_DARK
    p3.space_after = Pt(6)

    members = [
        ("GRS Akshitha", "AM.SC.U4CSE24221"),
        ("Sai Mohith Konduru", "AM.SC.U4CSE24230"),
        ("Nandini Kuchi", "AM.SC.U4CSE24231"),
        ("Leela Nayan P Kanjirapully", "AM.SC.U4CSE24232"),
        ("Mannsi Pentela", "AM.SC.U4CSE24247")
    ]
    for name, roll in members:
        pm = tf1.add_paragraph()
        pm.space_after = Pt(3)
        rn = pm.add_run()
        rn.text = f"•  {name}  "
        rn.font.size = Pt(13)
        rn.font.bold = True
        rn.font.color.rgb = COLOR_NAVY
        rr = pm.add_run()
        rr.text = f"({roll})"
        rr.font.size = Pt(12)
        rr.font.color.rgb = COLOR_MUTED

    # ==========================================
    # SLIDE 2: 01 • BACKGROUND
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, 2, 13, "01", "Background", "Problem statement — Why ambulances need signal priority at junctions")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "CRITICAL TIME LOST",
        [
            ("4-Way Intersection", "Congested urban junctions spanning North, East, South, and West."),
            ("The Golden Hour", "In acute cardiac and trauma cases, every 60 seconds of transit delay reduces patient survival probability by up to 10%."),
            ("Rigid Fixed Timers", "Traditional traffic controllers follow fixed timers without any perception of approaching emergency vehicles."),
            ("Dangerous Evasion", "Drivers often execute hazardous maneuvers into oncoming lanes to clear a path for trapped ambulances.")
        ],
        title_color=COLOR_RED, tag="The Challenge"
    )

    add_card(s2, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "THE ISSUE & PROJECT AIM",
        [
            ("Collision Risk", "Abruptly overriding signals creates severe broadside collision risks. Priority MUST be granted with controlled caution buffers."),
            ("Civilian Starvation", "Continuous emergency overrides along busy corridors can hold non-emergency lanes at permanent red lights."),
            ("Project Objective", "Automatically grant safe traffic-signal priority to approaching emergency vehicles via hardware interrupts."),
            ("Core Requirement", "Detect the ambulance, safely transition conflicting signals through yellow and all-red, sound warning alerts, and restore regular traffic fairly.")
        ],
        title_color=COLOR_BLUE, tag="The Solution"
    )

    # ==========================================
    # SLIDE 3: 02 • PROJECT PLAN
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, 3, 13, "02", "Project Plan", "Objectives — The key functions planned for the complete embedded system")

    # 6 Objective Cards (2 rows x 3 columns)
    cw = Inches(3.64)
    ch = Inches(2.35)
    r1_top = Inches(1.85)
    r2_top = Inches(4.5)

    add_card(s3, Inches(0.8), r1_top, cw, ch, "01 • DETECT", [
        ("Instant Sensing", "Sense approaching emergency vehicles from North, East, South, and West using pushbuttons on GPIO External Interrupts (EXTI).")
    ], title_color=COLOR_BLUE)

    add_card(s3, Inches(4.84), r1_top, cw, ch, "02 • TRANSITION", [
        ("Controlled Braking", "Safely change conflicting green signals to yellow (2 s), followed by all-red (1 s), before granting emergency green.")
    ], title_color=COLOR_AMBER)

    add_card(s3, Inches(8.88), r1_top, cw, ch, "03 • CLEAR", [
        ("Intersection Safety", "Enforce a mandatory 1-second all-red clearance interval before and after every green phase so conflicting lanes never cross.")
    ], title_color=COLOR_RED)

    add_card(s3, Inches(0.8), r2_top, cw, ch, "04 • ANTI-STARVATION", [
        ("Civilian Fairness", "Limit consecutive ambulance runs to 2; then force a full normal green phase for whichever lane waited the longest.")
    ], title_color=COLOR_GREEN)

    add_card(s3, Inches(4.84), r2_top, cw, ch, "05 • WARN & LOG", [
        ("Multi-Modal Alert", "Drive an audible warning buzzer on PA7 during priority and stream real-time diagnostic state telemetry over USART2.")
    ], title_color=COLOR_NAVY)

    add_card(s3, Inches(8.88), r2_top, cw, ch, "06 • QUEUE", [
        ("Zero Request Loss", "Store concurrent or repeated emergency requests in memory and service them fairly in round-robin order.")
    ], title_color=COLOR_BLUE)

    # ==========================================
    # SLIDE 4: 03 • SYSTEM DESIGN
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, 4, 13, "03", "System Design", "System Block Diagram — Input stage → Controller & timing stage → Output stage")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0),
        "INPUT & SENSING STAGE",
        [
            ("North Ambulance", "Pushbutton on PA0 (Falling edge EXTI Line 0)"),
            ("East Ambulance", "Pushbutton on PA1 (Falling edge EXTI Line 1)"),
            ("South Ambulance", "Pushbutton on PA4 (Falling edge EXTI Line 4)"),
            ("West Ambulance", "Pushbutton on PA8 (Falling edge EXTI Line 8)"),
            ("Internal Pull-Ups", "Normally 3.3V (HIGH); pressing grounds the input to 0V (LOW)."),
            ("EXTI + NVIC", "Asynchronous edge trigger immediately halts CPU wait-states.")
        ],
        title_color=COLOR_BLUE, tag="Inputs"
    )

    add_card(s4, Inches(4.8), Inches(1.8), Inches(3.7), Inches(5.0),
        "EMBEDDED CONTROLLER",
        [
            ("Target MCU", "STM32C031C6 (ARM Cortex-M0+ running at 12 MHz via HSI/4)."),
            ("Asynchronous ISR", "Captures direction in < 2 µs and applies a 50 ms contact debounce filter."),
            ("Non-Blocking FSM", "State machine driven by HAL_GetTick() without frozen delays."),
            ("Safety Interlock", "Hardware SetAllRed() driver prevents overlapping greens."),
            ("Fairness Engine", "Enforces 2-ambulance maximum and tracks lane wait times.")
        ],
        title_color=COLOR_NAVY, tag="Processing"
    )

    add_card(s4, Inches(8.9), Inches(1.8), Inches(3.6), Inches(5.0),
        "OUTPUT & ACTUATION STAGE",
        [
            ("Traffic LEDs (12 Pins)", "North: PB0/1/2\nEast: PB3/4/5\nSouth: PB6/7/8\nWest: PB9/PC6/PC7"),
            ("Current Limiting", "220 Ω resistors ensure safe LED operating current."),
            ("Audible Alert", "Piezoelectric warning buzzer on PA7 sounds during emergency green & yellow."),
            ("Telemetry Stream", "USART2 TX on PA2 (115200 baud) prints real-time events.")
        ],
        title_color=COLOR_GREEN, tag="Outputs"
    )

    # ==========================================
    # SLIDE 5: 04 • IMPLEMENTATION
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, 5, 13, "04", "Implementation", "Hardware components, pin mapping, and software development toolchain")

    add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "HARDWARE COMPONENTS & PINOUT",
        [
            ("Controller", "STM32C031C6 (ARM Cortex-M0+, 32 KB Flash, 12 KB SRAM)."),
            ("Traffic Signals", "12 active-high LEDs (Red, Yellow, Green for 4 approaches)."),
            ("Emergency Inputs", "4 pushbuttons with internal pull-ups on PA0, PA1, PA4, PA8."),
            ("Warning Alert", "Active buzzer driven by GPIO push-pull on PA7."),
            ("Current Limiting", "220 Ω resistors on all 12 LED channels."),
            ("Serial TX / RX", "PA2 (USART2_TX) and PA3 (USART2_RX) at 115200 baud.")
        ],
        title_color=COLOR_DARK, tag="Circuit Layout"
    )

    add_card(s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "SOFTWARE TOOLS & WORKFLOW",
        [
            ("STM32CubeMX", "Visual chip configuration, pinout assignment, clock tree generation (HSI DIV4 = 12 MHz), and HAL startup code."),
            ("STM32CubeIDE", "Embedded C development, GNU ARM GCC toolchain, compilation, linker script mapping, and ELF binary generation."),
            ("VS Code + Wokwi", "Cycle-accurate virtual hardware simulation, interactive pushbutton clicks, LED rendering, and live Serial Monitor streaming."),
            ("HAL & Non-Blocking", "Uses STM32C0 HAL with HAL_GetTick() delta timing, eliminating processor-freezing HAL_Delay() loops.")
        ],
        title_color=COLOR_BLUE, tag="Toolchain"
    )

    # ==========================================
    # SLIDE 6: 05 • CONTROL LOGIC
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, 6, 13, "05", "Control Logic", "Controller State Machine — 6 Non-blocking phases governing the junction")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "NORMAL TRAFFIC CYCLE",
        [
            ("Phase 1 • Normal Green (5 s)", "Active approach receives green; remaining 3 directions held red. Continuously scans for emergency requests."),
            ("Phase 2 • Normal Yellow (2 s)", "Caution clearance phase; warns approaching cars to stop smoothly when time expires or an emergency occurs."),
            ("Phase 3 • All-Red Clearance (1 s)", "All 4 approaches display red. Junction completely clears before switching to next direction."),
            ("Autonomous Sequence", "Rotates predictably: North → East → South → West → North.")
        ],
        title_color=COLOR_GREEN, tag="Regular Flow"
    )

    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "EMERGENCY PRIORITY PHASES",
        [
            ("Phase 4 • Ambulance Green (10 s)", "Requested approach gets green; buzzer turns ON at PA7; civilian traffic safely held red."),
            ("Phase 5 • Ambulance Yellow (2 s)", "Caution transition as ambulance clears; buzzer remains ON to warn nearby drivers."),
            ("Phase 6 • Ambulance All-Red (1 s)", "All red; buzzer turns OFF; scheduler checks whether another queued ambulance is waiting or if anti-starvation applies."),
            ("Safety Guarantee", "Direct green-to-green transitions are impossible — confirmed red required before any green turns on.")
        ],
        title_color=COLOR_RED, tag="Emergency Flow"
    )

    # ==========================================
    # SLIDE 7: 06 • EMBEDDED ENGINEERING
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, 7, 13, "06", "Embedded Engineering", "Interrupt callback resolution, contact debouncing, and atomic bitmask latching")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "THE EXTI CALLBACK RESOLUTION",
        [
            ("The Challenge", "In older STM32 families (F1/F4), a single HAL_GPIO_EXTI_Callback() handled all pin interrupts. In STM32C0, the HAL split this into separate Rising and Falling callbacks."),
            ("The Discarded Symbol", "Because our initial code implemented the legacy callback name, the STM32C0 linker discarded it as dead code, silently ignoring button presses."),
            ("The Solution", "Implementing HAL_GPIO_EXTI_Falling_Callback() properly overrode the weak HAL driver function, allowing button presses to trigger instantly."),
            ("Microsecond Response", "The ISR completes in < 2 µs by avoiding all blocking delays, logging, or LED operations inside interrupt context.")
        ],
        title_color=COLOR_BLUE, tag="HAL Integration"
    )

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "DEBOUNCE & BITMASK QUEUEING",
        [
            ("Mechanical Chatter", "Physical switches oscillate for milliseconds during closure, risking multiple false triggers per click."),
            ("50 ms Software Debounce", "Independent lastPressTime[4] timestamp guard ensures that subsequent transitions within 50 ms are safely suppressed."),
            ("Non-Destructive Bitmask", "Valid requests set bits in pendingRequests |= (1U << dir). Duplicate clicks on the same button do not inflate the queue."),
            ("Safe Acceptance", "A request bit is cleared only when the controller officially accepts that ambulance into service.")
        ],
        title_color=COLOR_DARK, tag="Signal Conditioning"
    )

    # ==========================================
    # SLIDE 8: 07 • FAIRNESS & ANTI-STARVATION
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, 8, 13, "07", "Fairness & Anti-Starvation", "Preventing civilian traffic starvation during heavy consecutive emergency calls")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "THE STARVATION THREAT",
        [
            ("Hospital Corridor Bottleneck", "Intersections adjacent to medical centers frequently experience multiple emergency vehicles arriving in close succession."),
            ("Civilian Gridlock", "Naive priority systems keep extending emergency green, trapping civilian lanes at red for 5-10 minutes."),
            ("Secondary Hazards", "Frustrated drivers run red lights; blocked intersections create secondary gridlock that traps future ambulances."),
            ("System Requirement", "An embedded system must be both responsive to emergencies and fundamentally fair to general traffic.")
        ],
        title_color=COLOR_RED, tag="Urban Traffic Dilemma"
    )

    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "THE DYNAMIC FAIRNESS ENGINE",
        [
            ("Configurable Limit", "MAX_CONSECUTIVE_AMBULANCES = 2. Tracks consecutive emergency runs without civilian traffic."),
            ("Wait-Time Tracking", "Array lastNormalGreenTime[4] records the timestamp of when each lane last enjoyed a normal green phase."),
            ("Guaranteed Civilian Phase", "When the 2nd ambulance finishes, the system forces a full 5-second normal green for the longest-waiting lane."),
            ("Uninterruptible Window", "This anti-starvation phase cannot be cut short by new emergency requests."),
            ("Zero Request Loss", "Waiting ambulance requests remain safely queued in memory and resume immediately after the civilian phase finishes.")
        ],
        title_color=COLOR_GREEN, tag="Scheduler Architecture"
    )

    # ==========================================
    # SLIDE 9: 08 • METHODOLOGY
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, 9, 13, "08", "Methodology", "Operating sequence — Step-by-step lifecycle of the junction controller")

    # 6 Steps in horizontal layout or 2 columns
    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "DETECTION & PREEMPTION SEQUENCE",
        [
            ("Step 1 • Normal Operation", "Junction rotates N → E → S → W using 5 s green, 2 s yellow, and 1 s all-red phases."),
            ("Step 2 • Request Ingest", "Ambulance button pulls pin to GND; EXTI triggers; ISR checks 50 ms debounce and sets direction bitmask."),
            ("Step 3 • Controlled Braking", "Conflicting active green transitions to yellow (2 s) and all-red (1 s) so cars clear the intersection safely."),
            ("Step 4 • Priority Phase", "Requested direction receives 10 seconds of green; warning buzzer sounds continuously on PA7.")
        ],
        title_color=COLOR_BLUE, tag="Steps 1 - 4"
    )

    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "CLEARANCE & FAIRNESS SEQUENCE",
        [
            ("Step 5 • Exit Caution", "Ambulance approach transitions to yellow (2 s) and all-red (1 s); warning buzzer shuts off."),
            ("Step 6 • Anti-Starvation Check", "If 2 consecutive ambulances have been served, grant a guaranteed 5 s normal green to the longest-waiting lane."),
            ("Step 7 • Queue Service", "If fewer than 2 ambulances or after civilian phase, service next pending emergency in round-robin sequence."),
            ("Step 8 • Return to Normal", "When queue is empty, resume normal rotation from the next approach in sequence.")
        ],
        title_color=COLOR_NAVY, tag="Steps 5 - 8"
    )

    # ==========================================
    # SLIDE 10: 09 • SIMULATION & VALIDATION
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, 10, 13, "09", "Simulation & Validation", "Observable testing scenarios and real-time UART telemetry in Wokwi")

    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "TEST SCENARIOS VALIDATED",
        [
            ("Test 1 • Normal Cycle", "Verified hands-free rotation through North, East, South, and West with exact 5s / 2s / 1s timing."),
            ("Test 2 • Emergency Preemption", "Pressed 'E' during North green. North safely transitioned to yellow (2s), all-red (1s), and East green lit for 10s with buzzer ON."),
            ("Test 3 • Concurrent Queueing", "Pressed 'S' and 'W' in rapid succession. South served first; West served immediately after without losing requests."),
            ("Test 4 • Anti-Starvation Test", "Triggered 'N', 'E', 'S'. First two served, then system forced a normal green for civilian traffic, then served 'S'."),
            ("Safety Check", "Zero overlapping greens across all tests; buzzer synchronized perfectly with priority.")
        ],
        title_color=COLOR_GREEN, tag="Wokwi Test Bench"
    )

    add_card(s10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "LIVE SERIAL TELEMETRY (USART2)",
        [
            ("System Boot", "[BOOT] Smart Ambulance Traffic Control Initialized"),
            ("Normal Cycle", "[STATE] Normal traffic: NORTH GREEN (5000 ms)"),
            ("Transition", "[STATE] Transition: NORTH YELLOW -> ALL RED"),
            ("Emergency Detect", "[EVENT] Ambulance request received: EAST"),
            ("Priority Grant", "[STATE] Ambulance priority granted: EAST (Consecutive: 1/2)"),
            ("Anti-Starvation", "[ALERT] Anti-starvation triggered! Max 2 consecutive ambulances reached."),
            ("Guaranteed Phase", "[STATE] Granting guaranteed normal green to longest-waiting: SOUTH"),
            ("Baud Rate", "115200 baud, 8N1 formatted output.")
        ],
        title_color=COLOR_DARK, tag="Console Logs"
    )

    # ==========================================
    # SLIDE 11: 10 • SDG ALIGNMENT
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, 11, 13, "10", "SDG Alignment", "United Nations Sustainable Development Goals — Health & Urban Mobility")

    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "SDG 3: GOOD HEALTH & WELL-BEING",
        [
            ("Target 3.6", "Halving global deaths and injuries from road traffic accidents."),
            ("Target 3.d", "Strengthening early warning, risk reduction, and rapid emergency response."),
            ("Preserving the Golden Hour", "Reducing intersection transit delays directly improves survival rates in cardiac arrests, acute trauma, and stroke emergencies."),
            ("Safe Transport", "Eliminates dangerous, unpredictable intersection crossings, protecting both patients and paramedics.")
        ],
        title_color=COLOR_GREEN, tag="UN SDG 3"
    )

    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "SDG 11: SUSTAINABLE CITIES & COMMUNITIES",
        [
            ("Target 11.2", "Providing access to safe, affordable, accessible, and sustainable transport systems for all."),
            ("Urban Congestion Relief", "Anti-starvation logic ensures regular commuter lanes and public transit are never permanently paralyzed by emergency traffic."),
            ("Safe & Resilient Infrastructure", "Mandatory yellow and all-red clearance intervals eliminate dangerous intersection collision risks."),
            ("Reduced Carbon Emissions", "Minimizing deadlocked traffic reduces unnecessary vehicle idling and localized greenhouse gas emissions.")
        ],
        title_color=COLOR_BLUE, tag="UN SDG 11"
    )

    # ==========================================
    # SLIDE 12: 11 • APPLICATIONS & FUTURE SCOPE
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, 12, 13, "11", "Applications & Future Scope", "Real-world smart city deployment and hardware sensor expansion")

    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "POTENTIAL APPLICATIONS",
        [
            ("Hospital Emergency Corridors", "Dynamic signal management along primary arterial roads connecting suburban highways to major trauma hospitals."),
            ("Intelligent Transportation Systems (ITS)", "Integration with centralized municipal traffic management centers for coordinated green-wave corridors."),
            ("Multi-Agency Emergency Preemption", "Easily expandable to fire rescue engines, disaster response vehicles, and police convoys with multi-tier priority levels.")
        ],
        title_color=COLOR_DARK, tag="Deployment"
    )

    add_card(s12, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "FUTURE HARDWARE EXPANSIONS",
        [
            ("V2X Wireless Transponders", "Replacing pushbuttons with DSRC / C-V2X (Vehicle-to-Infrastructure) radio beacons transmitting emergency authentication."),
            ("GPS Geo-Fencing", "Automatically preempting intersection signals when an approaching ambulance enters a 500-meter threshold."),
            ("Acoustic Siren Sensors", "Directional microphone arrays tuned to emergency siren frequencies (750 Hz - 1.5 kHz) for automated detection.")
        ],
        title_color=COLOR_BLUE, tag="Upgrades"
    )

    # ==========================================
    # SLIDE 13: 12 • CLOSING
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, 13, 13, "12", "Closing", "Project summary, key achievements, and technical references")

    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "SUMMARY OF ACHIEVEMENTS",
        [
            ("Complete Embedded Firmware", "Designed and verified a non-blocking state machine controller on the STM32C031C6 ARM Cortex-M0+ MCU."),
            ("EXTI Callback Fix", "Diagnosed and resolved the HAL_GPIO_EXTI_Falling_Callback linkage nuance in STM32C0."),
            ("Safety Interlock", "Hardware SetAllRed() driver physically eliminates conflicting green signals."),
            ("Anti-Starvation Engine", "Guarantees fair civilian traffic movement even during continuous emergency calls."),
            ("Virtual Prototyping", "Fully verified across 4 test scenarios using Wokwi simulation with live UART telemetry.")
        ],
        title_color=COLOR_BLUE, tag="Final Summary"
    )

    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0),
        "TECHNICAL REFERENCES",
        [
            ("Ref 1", "STMicroelectronics, STM32C031x4/x6/x8 Reference Manual (RM0490)."),
            ("Ref 2", "STMicroelectronics, STM32C031C6 Datasheet & STM32CubeC0 HAL Driver Documentation."),
            ("Ref 3", "IEEE, 'GPS and ZigBee Based Traffic Signal Preemption for Emergency Vehicles', ICICT, 2016."),
            ("Ref 4", "United Nations, Sustainable Development Goal 3: Good Health and Well-Being."),
            ("Ref 5", "United Nations, Sustainable Development Goal 11: Sustainable Cities and Communities.")
        ],
        title_color=COLOR_DARK, tag="Citations"
    )

    # Save presentation
    output_filename = "Smart_Ambulance_Traffic_System_Presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully saved to {output_filename}")

if __name__ == "__main__":
    create_presentation()
